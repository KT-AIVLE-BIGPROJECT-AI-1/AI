import json, uuid
from datetime import datetime, timezone
from pathlib import Path
from typing import Dict, Any

import requests
import tempfile
from pathlib import Path


import pdfplumber
import openpyxl
from docx import Document

from app.core.case_context import build_case_context   # 이미 있는 함수 활용

def download_presigned_to_temp(url: str) -> Path:
    """
    presigned_url → 임시 로컬 파일로 저장
    """
    response = requests.get(url)
    response.raise_for_status()

    suffix = Path(url).suffix or ".pdf"

    tmp = tempfile.NamedTemporaryFile(delete=False, suffix=suffix)
    tmp.write(response.content)
    tmp.close()

    return Path(tmp.name)
    
class DocumentAnalyzerV3:

    FORENSIC_PROMPT = """당신은 수사 보조 분석 담당자입니다.

아래에는 사건 개요와 문서 원문이 제공됩니다.

문서에 기재된 모든 객관적 정보를 빠짐없이 추출하여 설명문을 작성하되,
반드시 다음을 포함하십시오:

1. 문서의 종류 및 작성 목적
2. 문서에 기재된 주요 사실 (모든 객관적 정보 포함)
- 인적 사항 (이름, 생년월일, 주소, 연락처, 직업 등)
- 날짜/시간 정보
- 장소 정보
- 물품/증거물 정보
- 금액/거래 정보
- 기타 측정 가능한 수치

절대 규칙:
- 문서에 명시된 정보만 사용
- 추정·해석·법적 판단 금지
- 목록 사용 금지
- 연속된 설명문으로 6~10문장 작성

[사건 정보]
"""

    def __init__(self, client, model="gpt-4o-mini", temperature=0.0, chunk_size=3000):
        self.client = client
        self.model = model
        self.temperature = temperature
        self.chunk_size = chunk_size



    # ---------------- 문서 추출 ----------------
    def extract(self, path: Path) -> str:
        ext = path.suffix.lower()

        if ext == ".pdf":
            texts = []
            with pdfplumber.open(path) as pdf:
                for i, page in enumerate(pdf.pages, 1):
                    t = page.extract_text()
                    if t:
                        texts.append(f"[Page {i}]\n{t}")
            return "\n\n".join(texts)

        if ext == ".docx":
            doc = Document(path)
            lines = [p.text for p in doc.paragraphs if p.text.strip()]
            return "\n".join(lines)

        if ext in [".xlsx", ".xls"]:
            wb = openpyxl.load_workbook(path, data_only=True)
            blocks = []
            for ws in wb.worksheets:
                blocks.append(f"[Sheet: {ws.title}]")
                for row in ws.iter_rows():
                    values = [str(cell.value) for cell in row if cell.value]
                    if values:
                        blocks.append(" | ".join(values))
            return "\n".join(blocks)

        if ext in [".txt", ".csv", ".log"]:
            return path.read_text(encoding="utf-8")

        raise ValueError(f"지원하지 않는 문서 형식: {ext}")

    # ---------------- 청크 분할 ----------------
    def split_chunks(self, text: str):
        size = self.chunk_size
        overlap = int(size * 0.1)
        chunks = []
        start = 0
        while start < len(text):
            end = start + size
            chunks.append(text[start:end])
            start = end - overlap
        return chunks

    # ---------------- 분석 ----------------
    def analyze_one(self, text: str, case_json: dict = None) -> str:
        case_context = build_case_context(case_json, self.client)
        chunks = self.split_chunks(text)

        partials = []
        for chunk in chunks:
            prompt = self.FORENSIC_PROMPT + case_context + "\n[문서 원문]\n" + chunk

            resp = self.client.chat.completions.create(
                model=self.model,
                messages=[{"role": "user", "content": prompt}],
                temperature=self.temperature
            )

            partials.append(resp.choices[0].message.content.strip())

        return "\n".join(partials)


def analyze_document_to_fixed_json(analyzer, file_id, file_path, case_json, client):

    analysis_id = str(uuid.uuid4())
    created_at = datetime.now(timezone.utc).isoformat()

    try:
        # 👉 핵심 수정: presigned_url 다운로드
        local_path = download_presigned_to_temp(file_path)

        text = analyzer.extract(local_path)
        analysis_text = analyzer.analyze_one(text, case_json)

        status = "COMPLETED"

    except Exception as e:
        analysis_text = f"분석 실패: {str(e)}"
        status = "PENDING"

    return {
        "analysis_id": analysis_id,
        "file_id": file_id,
        "analysis_type": "DOC",
        "result_data": analysis_text,
        "created_at": created_at
    }
