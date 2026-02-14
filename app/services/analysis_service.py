# app/services/analysis_service.py
# 문서/이미지 분석 엔진 (내부 전용 - investigation_file_service에서 호출)

import base64
import openpyxl
import pdfplumber
from docx import Document
from PIL import Image
from datetime import datetime, timedelta
from pathlib import Path
from typing import List
from openai import OpenAI

from app.config import settings

client = OpenAI(api_key=settings.OPENAI_API_KEY)


# ============================================================
# 문서 분석 엔진
# ============================================================

class DocumentAnalyzerV3:

    FORENSIC_PROMPT = """당신은 수사 보조 분석 담당자입니다.

아래에는 사건 개요와 문서 원문이 제공됩니다.

문서에 기재된 모든 객관적 정보를 빠짐없이 추출하여 설명문을 작성하되,
반드시 다음을 포함하십시오:

1. 문서의 종류 및 작성 목적
2. 문서에 기재된 주요 사실 (모든 객관적 정보 포함)
   - 인적 사항 (이름, 생년월일, 주소, 연락처, 직업 등)
   - 날짜/시간 정보 (사건 발생 일시, 작성 일시, 기타 시간 정보)
   - 장소 정보 (주소, 위치, 장소 설명)
   - 물품/증거물 정보 (종류, 수량, 상태, 특징)
   - 상해/손상 정보 (부위, 정도, 치료 내용, 전치 기간)
   - 금액/거래 정보 (가격, 금액, 거래 내역)
   - 차량/물건 식별 정보 (번호판, 시리얼, IMEI 등)
   - 기타 측정 가능한 수치 (크기, 무게, 거리 등)
3. 사건 개요와 직접적으로 관련되는 부분

절대 규칙:
- 문서에 명시된 모든 구체적 정보를 빠짐없이 포함할 것
- 숫자, 날짜, 이름, 주소 등은 정확히 그대로 기록
- 문서에 없는 사실 추가 금지 (절대로 생성하지 마십시오)
- 추정, 해석, 판단, 법적 평가 금지
- 사건 개요에 없는 내용을 새로 만들지 말 것
- 목록/번호 사용 금지
- 설명문으로 작성 (6~10문장)

[사건 정보]
{case_info}

[문서 원문]
{document_text}
"""

    def __init__(self, client, model: str, temperature: float, chunk_size: int):
        self.client = client
        self.model = model
        self.temperature = temperature
        self.chunk_size = chunk_size

    def extract(self, path: Path) -> str:
        ext = path.suffix.lower()

        if ext == ".pdf":
            texts = []
            with pdfplumber.open(path) as pdf:
                for i, page in enumerate(pdf.pages, 1):
                    t = page.extract_text()
                    if t:
                        texts.append(f"[Page {i}]\n{t}")
            return "\n\n".join(texts)

        if ext == ".docx":
            doc = Document(path)
            lines = [p.text for p in doc.paragraphs if p.text.strip()]
            for table in doc.tables:
                for row in table.rows:
                    lines.append(" | ".join(cell.text for cell in row.cells))
            return "\n".join(lines)

        if ext in [".xlsx", ".xls"]:
            wb = openpyxl.load_workbook(path, data_only=True)
            blocks = []
            for ws in wb.worksheets:
                blocks.append(f"[Sheet: {ws.title}]")
                for row in ws.iter_rows():
                    values = []
                    for cell in row:
                        if cell.value is None:
                            continue
                        if isinstance(cell.value, (int, float)) and 30000 < cell.value < 60000:
                            base_dt = datetime(1899, 12, 30)
                            dt = base_dt + timedelta(days=int(cell.value))
                            values.append(dt.strftime("%Y-%m-%d"))
                        else:
                            values.append(str(cell.value))
                    if values:
                        blocks.append(" | ".join(values))
            return "\n".join(blocks)

        if ext in [".txt", ".csv", ".log"]:
            for enc in ["utf-8", "cp949", "euc-kr", "latin1"]:
                try:
                    return path.read_text(encoding=enc)
                except Exception:
                    continue
            raise ValueError("텍스트 인코딩 판독 실패")

        raise ValueError(f"지원하지 않는 문서 형식: {ext}")

    def split_chunks(self, text: str) -> List[str]:
        size = self.chunk_size
        overlap = int(size * 0.1)
        chunks = []
        start = 0
        while start < len(text):
            end = start + size
            chunks.append(text[start:end])
            start = end - overlap
        return chunks

    def _format_case_info(self, case_json: dict) -> str:
        if not case_json:
            return "사건 정보 없음"
        info_parts = []
        if "title" in case_json:
            info_parts.append(f"사건명: {case_json['title']}")
        if "type" in case_json:
            info_parts.append(f"사건 유형: {case_json['type']}")
        if "description" in case_json:
            info_parts.append(f"사건 개요: {case_json['description']}")
        occurred_time = case_json.get('occurred_at', case_json.get('occured_at', ''))
        if occurred_time:
            info_parts.append(f"발생 일시: {occurred_time}")
        if "location" in case_json:
            info_parts.append(f"발생 장소: {case_json['location']}")
        if "summary" in case_json and case_json["summary"]:
            summary = case_json["summary"]
            if isinstance(summary, dict) and any(summary.values()):
                info_parts.append("\n사건 요약:")
                for key, value in summary.items():
                    if value:
                        info_parts.append(f"  - {key}: {value}")
        return "\n".join(info_parts)

    def analyze_one(self, text: str, case_json: dict = None, client=None) -> str:
        api_client = client or self.client
        case_info = self._format_case_info(case_json)
        chunks = self.split_chunks(text)
        partials = []
        for chunk in chunks:
            prompt = self.FORENSIC_PROMPT.format(case_info=case_info, document_text=chunk)
            response = api_client.chat.completions.create(
                model=self.model,
                messages=[{"role": "user", "content": prompt}],
                temperature=self.temperature
            )
            partials.append(response.choices[0].message.content.strip())
        return "\n".join(partials)


# ============================================================
# 이미지 분석 엔진
# ============================================================

class ImageAnalyzerV2:

    FORENSIC_PROMPT = """당신은 수사 보조 이미지 분석 담당자입니다.

아래에는 사건 개요와 사진이 제공됩니다.

사진에 실제로 보이는 모든 객관적 정보를 빠짐없이 추출하여 설명문을 작성하되,
반드시 다음을 포함하십시오:

1. 사진에 나타난 모든 관찰 가능한 사실
   - 인물 (복장, 외형, 자세, 행동)
   - 물체/물품 (종류, 색상, 크기, 상태, 위치)
   - 장소 (건물, 도로, 표지판, 지형, 환경)
   - 시간 정보 (그림자, 조명, 날씨로 추정 가능한 시간대)
   - 텍스트 (간판, 표지판, 번호판 등 읽을 수 있는 모든 글자)
   - 거리/위치 관계 (사람 간 거리, 물체 배치)
2. 사건 개요와 직접적으로 연결되는 부분

절대 규칙:
- 사진에 보이는 모든 구체적 정보를 빠짐없이 포함할 것
- 색상, 위치, 크기 등은 가능한 구체적으로 기술
- 사진에 보이지 않는 내용 추가 금지 (절대로 생성하지 마십시오)
- 사건 개요에 없는 내용을 새로 만들지 말 것
- 추정, 단정, 법적 판단 금지
- 미학적 묘사 금지
- 목록/번호 사용 금지
- 설명문으로 작성 (6~10문장)

[사건 정보]
{case_info}

[사진 메타]
{meta_text}
"""

    def __init__(self, client, model: str):
        self.client = client
        self.model = model

    def _image_to_data_url(self, path: Path) -> str:
        with open(path, "rb") as f:
            b64 = base64.b64encode(f.read()).decode("utf-8")
        ext = path.suffix.replace(".", "").lower()
        if ext == "jpg":
            ext = "jpeg"
        return f"data:image/{ext};base64,{b64}"

    def _get_image_meta(self, path: Path) -> str:
        with Image.open(path) as im:
            w, h = im.size
        return f"- 파일명: {path.name}\n- 해상도: {w}x{h}\n"

    def _format_case_info(self, case_json: dict) -> str:
        if not case_json:
            return "사건 정보 없음"
        info_parts = []
        if "title" in case_json:
            info_parts.append(f"사건명: {case_json['title']}")
        if "type" in case_json:
            info_parts.append(f"사건 유형: {case_json['type']}")
        if "description" in case_json:
            info_parts.append(f"사건 개요: {case_json['description']}")
        occurred_time = case_json.get('occurred_at', case_json.get('occured_at', ''))
        if occurred_time:
            info_parts.append(f"발생 일시: {occurred_time}")
        if "location" in case_json:
            info_parts.append(f"발생 장소: {case_json['location']}")
        if "summary" in case_json and case_json["summary"]:
            summary = case_json["summary"]
            if isinstance(summary, dict) and any(summary.values()):
                info_parts.append("\n사건 요약:")
                for key, value in summary.items():
                    if value:
                        info_parts.append(f"  - {key}: {value}")
        return "\n".join(info_parts)

    def analyze_one(self, image_path: Path, case_json: dict = None, client=None) -> str:
        api_client = client or self.client
        case_info = self._format_case_info(case_json)
        meta_text = self._get_image_meta(image_path)
        data_url = self._image_to_data_url(image_path)
        full_text = self.FORENSIC_PROMPT.format(case_info=case_info, meta_text=meta_text)
        response = api_client.chat.completions.create(
            model=self.model,
            messages=[{
                "role": "user",
                "content": [
                    {"type": "text", "text": full_text},
                    {"type": "image_url", "image_url": {"url": data_url}}
                ]
            }],
            temperature=0.0
        )
        return response.choices[0].message.content.strip()


# ============================================================
# 키워드 추출 유틸
# ============================================================

def extract_keywords(text: str, client_instance=None) -> List[str]:
    try:
        api_client = client_instance or client
        response = api_client.chat.completions.create(
            model="gpt-4o-mini",
            messages=[{
                "role": "user",
                "content": f"다음 수사 분석 텍스트에서 핵심 키워드 5~10개를 쉼표로 구분하여 출력하십시오. 키워드만 출력하고 다른 설명은 하지 마십시오.\n\n{text[:1500]}"
            }],
            temperature=0.1,
            max_tokens=200
        )
        raw = response.choices[0].message.content.strip()
        return [k.strip() for k in raw.split(",") if k.strip()]
    except Exception:
        return []


# ============================================================
# Analyzer 인스턴스 (investigation_file_service에서 import하여 사용)
# ============================================================

document_analyzer = DocumentAnalyzerV3(
    client=client, model="gpt-4.1", temperature=0.0, chunk_size=3000
)

image_analyzer = ImageAnalyzerV2(
    client=client, model="gpt-4.1"
)