# app/services/analysis_document_service.py

import uuid
import openpyxl
import pdfplumber
from docx import Document
from datetime import datetime, timedelta, timezone
from pathlib import Path
from typing import List, Dict, Any
from openai import OpenAI

from app.config import settings
from app.models.analysis_document_schema import (
    AnalysisDocumentRequest,
    AnalysisDocumentResponse
)

client = OpenAI(api_key=settings.OPENAI_API_KEY)


class DocumentAnalyzerV3:

    FORENSIC_PROMPT = """당신은 수사 보조 분석 담당자입니다.

아래에는 사건 개요와 문서 원문이 제공됩니다.

문서에 기재된 모든 객관적 정보를 빠짐없이 추출하여 설명문을 작성하되,
반드시 다음을 포함하십시오:

1. 문서의 종류 및 작성 목적
2. 문서에 기재된 주요 사실 (모든 객관적 정보 포함)
   - 인적 사항 (이름, 생년월일, 주소, 연락처, 직업 등)
   - 날짜/시간 정보 (사건 발생 일시, 작성 일시, 기타 시간 정보)
   - 장소 정보 (주소, 위치, 장소 설명)
   - 물품/증거물 정보 (종류, 수량, 상태, 특징)
   - 상해/손상 정보 (부위, 정도, 치료 내용, 전치 기간)
   - 금액/거래 정보 (가격, 금액, 거래 내역)
   - 차량/물건 식별 정보 (번호판, 시리얼, IMEI 등)
   - 기타 측정 가능한 수치 (크기, 무게, 거리 등)
3. 사건 개요와 직접적으로 관련되는 부분

절대 규칙:
- 문서에 명시된 모든 구체적 정보를 빠짐없이 포함할 것
- 숫자, 날짜, 이름, 주소 등은 정확히 그대로 기록
- 문서에 없는 사실 추가 금지
- 추정, 해석, 판단, 법적 평가 금지
- 사건 개요에 없는 내용을 새로 만들지 말 것
- 목록/번호 사용 금지
- 설명문으로 작성 (6~10문장)

[사건 정보]
{case_info}

[문서 원문]
{document_text}
"""

    def __init__(self, client, model: str, temperature: float, chunk_size: int):
        self.client = client
        self.model = model
        self.temperature = temperature
        self.chunk_size = chunk_size

    def extract(self, path: Path) -> str:
        ext = path.suffix.lower()

        if ext == ".pdf":
            texts = []
            with pdfplumber.open(path) as pdf:
                for i, page in enumerate(pdf.pages, 1):
                    t = page.extract_text()
                    if t:
                        texts.append(f"[Page {i}]\n{t}")
            return "\n\n".join(texts)

        if ext == ".docx":
            doc = Document(path)
            lines = [p.text for p in doc.paragraphs if p.text.strip()]
            for table in doc.tables:
                for row in table.rows:
                    lines.append(" | ".join(cell.text for cell in row.cells))
            return "\n".join(lines)

        if ext in [".xlsx", ".xls"]:
            wb = openpyxl.load_workbook(path, data_only=True)
            blocks = []
            for ws in wb.worksheets:
                blocks.append(f"[Sheet: {ws.title}]")
                for row in ws.iter_rows():
                    values = []
                    for cell in row:
                        if cell.value is None:
                            continue
                        if isinstance(cell.value, (int, float)) and 30000 < cell.value < 60000:
                            base = datetime(1899, 12, 30)
                            dt = base + timedelta(days=int(cell.value))
                            values.append(dt.strftime("%Y-%m-%d"))
                        else:
                            values.append(str(cell.value))
                    if values:
                        blocks.append(" | ".join(values))
            return "\n".join(blocks)

        if ext in [".txt", ".csv", ".log"]:
            for enc in ["utf-8", "cp949", "euc-kr", "latin1"]:
                try:
                    return path.read_text(encoding=enc)
                except Exception:
                    continue
            raise ValueError("텍스트 인코딩 판독 실패")

        raise ValueError(f"지원하지 않는 문서 형식: {ext}")

    def split_chunks(self, text: str) -> List[str]:
        size = self.chunk_size
        overlap = int(size * 0.1)
        chunks = []
        start = 0
        while start < len(text):
            end = start + size
            chunks.append(text[start:end])
            start = end - overlap
        return chunks

    def _format_case_info(self, case_json: dict) -> str:
        if not case_json:
            return "사건 정보 없음"
        
        info_parts = []
        if "title" in case_json:
            info_parts.append(f"사건명: {case_json['title']}")
        if "type" in case_json:
            info_parts.append(f"사건 유형: {case_json['type']}")
        if "description" in case_json:
            info_parts.append(f"사건 개요: {case_json['description']}")
        occurred_time = case_json.get('occurred_at', case_json.get('occured_at', ''))
        if occurred_time:
            info_parts.append(f"발생 일시: {occurred_time}")
        if "location" in case_json:
            info_parts.append(f"발생 장소: {case_json['location']}")
        if "summary" in case_json and case_json["summary"]:
            summary = case_json["summary"]
            if isinstance(summary, dict) and any(summary.values()):
                info_parts.append("\n사건 요약:")
                for key, value in summary.items():
                    if value:
                        info_parts.append(f"  - {key}: {value}")
        return "\n".join(info_parts)

    def analyze_one(self, text: str, case_json: dict = None, client=None) -> str:
        api_client = client or self.client
        case_info = self._format_case_info(case_json)
        chunks = self.split_chunks(text)
        partials = []

        for chunk in chunks:
            prompt = self.FORENSIC_PROMPT.format(
                case_info=case_info,
                document_text=chunk
            )
            response = api_client.chat.completions.create(
                model=self.model,
                messages=[{"role": "user", "content": prompt}],
                temperature=self.temperature
            )
            partials.append(response.choices[0].message.content.strip())

        return "\n".join(partials)


def analyze_document_to_fixed_json(
    analyzer: DocumentAnalyzerV3,
    analysis_id: str,
    file_id: str,
    file_path: Path,
    case_json: dict = None,
    client = None
) -> Dict[str, Any]:
    created_at = datetime.now(timezone.utc).isoformat()
    status = "PROCESSING"

    try:
        text = analyzer.extract(file_path)
        analysis_text = analyzer.analyze_one(text, case_json, client)
        status = "COMPLETED"
    except Exception as e:
        analysis_text = f"분석 실패: {str(e)}"
        status = "PENDING"

    return {
        "analysis_id": analysis_id,
        "file_id": file_id,
        "file_type": "DOCUMENT",
        "status": status,
        "result_data": analysis_text,
        "created_at": created_at
    }


document_analyzer = DocumentAnalyzerV3(
    client=client,
    model="gpt-4o",
    temperature=0.3,
    chunk_size=3000
)


class AnalysisDocumentService:
    @staticmethod
    def analyze_document(request: AnalysisDocumentRequest) -> AnalysisDocumentResponse:
        print(f"\n📄 문서 분석 시작: {request.file_id}")
        
        file_path = Path(request.file_path)
        analysis_id = str(uuid.uuid4())
        case_json = None
        
        result = analyze_document_to_fixed_json(
            analyzer=document_analyzer,
            analysis_id=analysis_id,
            file_id=request.file_id,
            file_path=file_path,
            case_json=case_json,
            client=client
        )
        
        print(f"✅ 문서 분석 완료: {result['status']}")
        
        from app.services.analysis_callback_service import AnalysisCallbackService
        AnalysisCallbackService.register_analysis_result(
            analysis_id=result["analysis_id"],
            result_data=result
        )
        
        return AnalysisDocumentResponse(**result)